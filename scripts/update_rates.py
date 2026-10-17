"""
Monthly NZ mortgage rate updater.

Downloads the Reserve Bank of New Zealand's mortgage rate series (B20
standard rates and B21 special rates), and writes into ./docs (served by
GitHub Pages):

    historical-mortgage-rates.xlsx   the spreadsheet for the website
    nz-mortgage-rates.json           data + headline figures for the app
    nz-mortgage-rates.csv            standard rates, plain data
    nz-mortgage-rates-b21.csv        special rates, plain data

Fallback: if RBNZ blocks the automated download, put hb20.xlsx / hb21.xlsx
into ./inbox and the script uses those instead (then removes them).

Safety checks stop the run, leaving the published files untouched, if the
new data looks wrong: missing terms, implausible rates, or fewer months than
are already published.
"""

from __future__ import annotations

import io
import json
import re
import sys
import urllib.request
from datetime import date, datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import workbook  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "docs"
INBOX = ROOT / "inbox"
LOGO = ROOT / "assets" / "logo.png"

XLSX_NAME = "historical-mortgage-rates.xlsx"
JSON_NAME = "nz-mortgage-rates.json"

# Bump this when the layout of the outputs changes, so the next run rebuilds
# them even if RBNZ hasn't released new data.
FORMAT_VERSION = 5

SERIES = {
    "B20": {
        "label": "Standard",
        "url": "https://www.rbnz.govt.nz/-/media/project/sites/rbnz/files/statistics/series/b/b20/hb20.xlsx",
        "page": "https://www.rbnz.govt.nz/statistics/series/exchange-and-interest-rates/new-residential-mortgage-standard-interest-rates",
        "file": "hb20.xlsx",
        # Months before this are ignored, so the whole sheet runs from
        # January 2005 (whole years, and every 10-year period starts in January).
        "start": "2005-01",
        "csv": "nz-mortgage-rates.csv",
    },
    "B21": {
        "label": "Special",
        "url": "https://www.rbnz.govt.nz/-/media/project/sites/rbnz/files/statistics/series/b/b21/hb21.xlsx",
        "page": "https://www.rbnz.govt.nz/statistics/series/exchange-and-interest-rates/new-residential-mortgage-special-interest-rates",
        "file": "hb21.xlsx",
        "start": "2017-01",
        "csv": "nz-mortgage-rates-b21.csv",
    },
}

# Terms in display order: (key, pattern matched against RBNZ column headings)
TERM_PATTERNS = [
    ("floating", r"\bfloat"),
    ("6m", r"\b6[\s-]*month|six[\s-]*month"),
    ("1y", r"\b1[\s-]*year|one[\s-]*year|\b12[\s-]*month"),
    ("18m", r"\b18[\s-]*month"),  # recognised so it isn't mistaken for another term; not published
    ("2y", r"\b2[\s-]*year|two[\s-]*year"),
    ("3y", r"\b3[\s-]*year|three[\s-]*year"),
    ("4y", r"\b4[\s-]*year|four[\s-]*year"),
    ("5y", r"\b5[\s-]*year|five[\s-]*year"),
]
PUBLISHED_TERMS = workbook.TERM_KEYS
# Terms that must be present. Special rates (B21) may not include floating.
REQUIRED_TERMS = {"B20": ["floating", "6m", "1y", "2y", "3y", "4y", "5y"],
                  "B21": ["1y", "2y", "3y", "5y"]}

USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/126.0 Safari/537.36"
)


class DataProblem(Exception):
    """Raised when the RBNZ data can't be fetched or doesn't look right."""


# ---------------------------------------------------------------- fetching

def fetch(series_code: str) -> tuple[bytes, str]:
    """Return (file bytes, where it came from). An inbox file wins over the download."""
    info = SERIES[series_code]
    local = INBOX / info["file"]
    if local.exists():
        return local.read_bytes(), f"inbox/{info['file']}"

    req = urllib.request.Request(info["url"], headers={
        "User-Agent": USER_AGENT,
        "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*",
        "Referer": info["page"],
    })
    try:
        with urllib.request.urlopen(req, timeout=60) as resp:
            data = resp.read()
    except Exception as exc:  # network error, 403 from bot protection, etc.
        raise DataProblem(
            f"Couldn't download {series_code} from RBNZ ({exc}).\n"
            f"Fallback: download {info['file']} yourself from {info['page']}\n"
            f"and upload it to the inbox/ folder of this repository."
        ) from exc

    # An .xlsx is a zip file, which always starts with 'PK'. Anything else is
    # usually a bot-check web page.
    if not data.startswith(b"PK"):
        raise DataProblem(
            f"RBNZ returned a web page instead of the {series_code} spreadsheet "
            f"(likely bot protection).\n"
            f"Fallback: download {info['file']} from {info['page']}\n"
            f"and upload it to the inbox/ folder of this repository."
        )
    return data, info["url"]


# ---------------------------------------------------------------- parsing

def _as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%b %Y", "%B %Y", "%Y-%m"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                pass
    return None


def _match_term(heading: str) -> str | None:
    text = heading.lower()
    for key, pattern in TERM_PATTERNS:
        if re.search(pattern, text):
            return key
    return None


def parse(data: bytes, series_code: str) -> dict[str, dict[str, float]]:
    """
    Return {"YYYY-MM": {"floating": 7.24, "1y": 6.1, ...}, ...} (percent).

    Finds rows whose first cell is a date (the data), then reads the heading
    cells above each column to work out which term it is. Checks every
    sheet, so it copes with RBNZ splitting terms across sheets.
    """
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    result: dict[str, dict[str, float]] = {}
    found: dict[str, str] = {}

    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        if not rows:
            continue
        first_data = next(
            (i for i, r in enumerate(rows) if r and _as_date(r[0]) is not None
             and any(isinstance(v, (int, float)) for v in r[1:])),
            None,
        )
        if first_data is None:
            continue

        width = max(len(r) for r in rows)
        header_rows = rows[:first_data]
        col_terms: dict[int, str] = {}
        for c in range(1, width):
            # Use the heading nearest the data that names a term; group labels
            # higher up (e.g. "terms greater than two years") are skipped.
            key, heading = None, ""
            for r in reversed(header_rows):
                if c < len(r) and r[c] is not None:
                    key = _match_term(str(r[c]))
                    if key:
                        heading = str(r[c])
                        break
            if key and key not in found:
                col_terms[c] = key
                found[key] = f"{ws.title}!{get_column_letter(c + 1)}: {heading.strip()[:80]}"

        for r in rows[first_data:]:
            d = _as_date(r[0]) if r else None
            if d is None:
                continue
            month = f"{d.year:04d}-{d.month:02d}"
            for c, key in col_terms.items():
                v = r[c] if c < len(r) else None
                if isinstance(v, (int, float)) and key in PUBLISHED_TERMS:
                    result.setdefault(month, {})[key] = round(float(v), 2)

    missing = [k for k in REQUIRED_TERMS[series_code] if k not in found]
    if missing:
        raise DataProblem(
            f"{series_code}: couldn't find columns for {missing}. RBNZ may have "
            f"changed the file layout. Columns found: {found or 'none'}"
        )
    print(f"{series_code} columns:")
    for key, where in found.items():
        print(f"  {key:9s} <- {where}")

    start = SERIES[series_code]["start"]
    return {m: v for m, v in result.items() if m >= start}


# ---------------------------------------------------------------- checks

def check(rows: dict[str, dict[str, float]], series_code: str) -> None:
    if len(rows) < 12:
        raise DataProblem(f"{series_code}: only {len(rows)} months found, expected years of history.")
    for month, terms in rows.items():
        for key, v in terms.items():
            if not 0 < v < 30:
                raise DataProblem(f"{series_code}: implausible rate {v}% for {key} in {month}.")

    previous = OUT_DIR / JSON_NAME
    if previous.exists():
        old = json.loads(previous.read_text())
        old_months = len(old.get("series", {}).get(series_code, {}).get("months", []))
        if len(rows) < old_months:
            raise DataProblem(
                f"{series_code}: new data has {len(rows)} months but {old_months} are "
                f"already published. Refusing to overwrite with less data."
            )

    latest = max(rows)
    y, m = map(int, latest.split("-"))
    today = date.today()
    age = (today.year - y) * 12 + (today.month - m)
    if age > 4:
        print(f"WARNING: latest {series_code} month is {latest}, {age} months old.")


# ---------------------------------------------------------------- outputs

def _pct(x, places=4):
    return None if x is None else round(x * 100, places)


def write_outputs(all_series: dict[str, dict], sources: dict[str, str]) -> bool:
    """Write all output files. Returns False if nothing needed updating."""
    OUT_DIR.mkdir(exist_ok=True)
    generated = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    standard, special = all_series["B20"], all_series.get("B21", {})
    figs = workbook.compute(standard, special)

    def terms_in(rows):
        present = {k for t in rows.values() for k in t}
        return [k for k in PUBLISHED_TERMS if k in present]

    # JSON for the app ---------------------------------------------------
    payload = {
        "format_version": FORMAT_VERSION,
        "source": "Reserve Bank of New Zealand",
        "generated": generated,
        "units": "percent per annum, end-of-month average",
        "summary": {
            "latest_month": figs["latest"],
            "standard": {k: {s: _pct(v) for s, v in d.items()}
                         for k, d in figs["standard_overall"].items() if d},
            "special": {k: {s: _pct(v) for s, v in d.items()}
                        for k, d in figs["special_overall"].items() if d},
            "standard_10yr_plus_average": {k: _pct(v) for k, v in figs["long_period_average"].items()},
            "cheapest_term_share": {k: _pct(v, 2) for k, v in figs["cheapest_share"].items()},
        },
        "series": {},
    }
    for code, rows in all_series.items():
        keys = terms_in(rows)
        months = sorted(rows)
        payload["series"][code] = {
            "name": f"New residential mortgage {SERIES[code]['label'].lower()} interest rates",
            "source_page": SERIES[code]["page"],
            "latest_month": months[-1],
            "terms": [{"key": k, "name": workbook.TERM_NAMES[k]} for k in keys],
            "months": [{"month": mo, **{k: rows[mo].get(k) for k in keys}} for mo in months],
        }

    # Skip if the data and layout haven't changed since the last publish, so
    # the repository only gets a commit when RBNZ releases new data.
    json_path = OUT_DIR / JSON_NAME
    if json_path.exists() and (OUT_DIR / XLSX_NAME).exists():
        old = json.loads(json_path.read_text())
        same_data = ({c: s["months"] for c, s in old.get("series", {}).items()} ==
                     {c: s["months"] for c, s in payload["series"].items()})
        if same_data and old.get("format_version") == FORMAT_VERSION:
            print("No new data since last publish – nothing to update.")
            return False

    # Spreadsheet ---------------------------------------------------------
    nz_date = datetime.now(ZoneInfo("Pacific/Auckland")).strftime("%-d %B %Y")
    workbook.build(
        OUT_DIR / XLSX_NAME, standard, special, figs, nz_date,
        logo_path=LOGO if LOGO.exists() else None,
        source_pages={c: SERIES[c]["page"] for c in all_series},
    )

    # CSV -----------------------------------------------------------------
    for code, rows in all_series.items():
        keys = terms_in(rows)
        out = ["Month," + ",".join(workbook.TERM_NAMES[k] for k in keys)]
        for mo in sorted(rows, reverse=True):
            out.append(mo + "," + ",".join(
                "" if rows[mo].get(k) is None else f"{rows[mo][k]:.2f}" for k in keys))
        (OUT_DIR / SERIES[code]["csv"]).write_text("\n".join(out) + "\n")

    json_path.write_text(json.dumps(payload, indent=1))
    return True


# ---------------------------------------------------------------- main

def main() -> int:
    all_series, sources = {}, {}
    try:
        for code in SERIES:
            data, source = fetch(code)
            rows = parse(data, code)
            check(rows, code)
            all_series[code] = rows
            sources[code] = source
            print(f"{code}: {len(rows)} months, {min(rows)} to {max(rows)} (from {source})")
    except DataProblem as problem:
        print(f"\nSTOPPED – published files were not changed.\n{problem}", file=sys.stderr)
        return 1

    if write_outputs(all_series, sources):
        print(f"Updated files in {OUT_DIR}")

    # A file dropped into inbox/ is used once, then removed, so next month's
    # run goes back to downloading from RBNZ automatically.
    for code, source in sources.items():
        if source.startswith("inbox/"):
            (INBOX / SERIES[code]["file"]).unlink()
            print(f"Used and removed {source}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
